"""Human Protein Atlas per-organ RNA and protein expression of the salivary
glandular-backbone genes. Self-contained: fetches HPA data (cached under data/),
writes the supplementary table (tables/) and the two-sided RNA-vs-protein figure
(figures/). Run from anywhere:  python hpa_expression.py

Data source: Human Protein Atlas (https://www.proteinatlas.org), version 23.
RNA = consensus tissue nTPM (search API); protein = antibody-based IHC score
(Not detected / Low / Medium / High), parsed from the per-gene XML. CC BY-SA 4.0.
"""
import os, gzip, json, time, urllib.request, xml.etree.ElementTree as ET
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
matplotlib.rcParams.update({"svg.fonttype": "none", "pdf.fonttype": 42,
                            "font.family": "sans-serif", "font.sans-serif": ["Helvetica", "Arial", "DejaVu Sans"]})
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Patch
from matplotlib.cm import ScalarMappable
from matplotlib.colors import Normalize
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = os.path.dirname(os.path.abspath(__file__))
DATA, FIG, TAB = f"{BASE}/data", f"{BASE}/figures", f"{BASE}/tables"
for d in (DATA, FIG, TAB): os.makedirs(d, exist_ok=True)

# gene -> Ensembl gene id (human)
ENS = {"AQP5": "ENSG00000161798", "ZG16B": "ENSG00000162078", "CRACR2A": "ENSG00000130038",
       "KCNN4": "ENSG00000104783", "CLDN10": "ENSG00000134873", "CGREF1": "ENSG00000138028",
       "WFDC12": "ENSG00000168703"}
GENES = list(ENS)
TISSUES = ["adipose tissue", "adrenal gland", "appendix", "bone marrow", "breast", "cerebral cortex",
           "cerebellum", "cervix", "colon", "duodenum", "endometrium", "epididymis", "esophagus",
           "fallopian tube", "gallbladder", "heart muscle", "kidney", "liver", "lung", "lymph node",
           "ovary", "pancreas", "parathyroid gland", "pituitary gland", "placenta", "prostate", "rectum",
           "retina", "salivary gland", "seminal vesicle", "skeletal muscle", "skin", "small intestine",
           "smooth muscle", "spleen", "stomach", "testis", "thymus", "thyroid gland", "tongue", "tonsil",
           "urinary bladder", "vagina"]
SAL = "salivary gland"
LEVELS = {"not detected": 0, "low": 1, "medium": 2, "high": 3}
LVLNAME = {0: "Not detected", 1: "Low", 2: "Medium", 3: "High"}


def _get(url):
    raw = urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60).read()
    try: raw = gzip.decompress(raw)
    except Exception: pass
    return raw


def fetch():
    """Download RNA (nTPM) and protein (IHC) per organ; cache to data/."""
    rna_csv, pro_csv = f"{DATA}/hpa_rna_tissue_nTPM.csv", f"{DATA}/hpa_protein_tissue_ihc.csv"
    if os.path.exists(rna_csv) and os.path.exists(pro_csv):
        return pd.read_csv(rna_csv), pd.read_csv(pro_csv)
    cols = "g,eg," + ",".join("t_RNA_" + t.replace(" ", "_") for t in TISSUES)
    rna_rows, pro_rows = [], []
    for sym, e in ENS.items():
        d = json.loads(_get(f"https://www.proteinatlas.org/api/search_download.php?search={sym}&format=json&columns={cols}"))
        r = [x for x in d if x.get("Ensembl") == e][0]
        for t in TISSUES:
            v = r.get(f"Tissue RNA - {t} [nTPM]")
            rna_rows.append({"gene": sym, "tissue": t, "nTPM": float(v) if v not in (None, "") else None})
        xml = f"{DATA}/{sym}.xml"
        if not os.path.exists(xml): open(xml, "wb").write(_get(f"https://www.proteinatlas.org/{e}.xml"))
        root = ET.parse(xml).getroot()
        for te in root.iter("tissueExpression"):
            if te.get("technology") == "IHC" and te.get("assayType") == "tissue":
                for data in te.iter("data"):
                    tis, lvl = data.find("tissue"), data.find("level")
                    if tis is not None and lvl is not None and lvl.get("type") == "expression":
                        pro_rows.append({"gene": sym, "tissue": tis.text, "level": lvl.text})
                break
        time.sleep(0.3)
    rna = pd.DataFrame(rna_rows); pro = pd.DataFrame(pro_rows).drop_duplicates(["gene", "tissue"])
    rna.to_csv(rna_csv, index=False); pro.to_csv(pro_csv, index=False)
    return rna, pro


def matrices(rna, pro):
    """Return per-organ RNA nTPM (R) and protein level 0-3 (P), organs x genes, salivary first."""
    R = rna.pivot(index="tissue", columns="gene", values="nTPM").reindex(columns=GENES)
    pro = pro.copy(); pro["lv"] = pro.level.str.lower().map(LEVELS)
    pro["tissue"] = pro.tissue.str.replace(r"\s+\d+$", "", regex=True).str.lower()
    P = pro.groupby(["tissue", "gene"])["lv"].max().unstack().reindex(columns=GENES).reindex(R.index)
    order = [SAL] + sorted([t for t in R.index if t != SAL], key=lambda t: -np.nan_to_num(R.loc[t].max()))
    return R.loc[order], P.loc[order]


def write_table(R, P):
    """Supplementary table: two sheets (RNA nTPM, protein IHC level), organs x genes."""
    wb = Workbook(); F = "Arial"
    hf, hfont = PatternFill("solid", fgColor="1F4E79"), Font(name=F, bold=True, color="FFFFFF")
    salfill = PatternFill("solid", fgColor="FCE4D6")

    def sheet(ws, mat, fmt):
        ws.append(["Organ"] + [f"{g} ({ENS[g]})" for g in GENES])
        for c in ws[1]: c.font = hfont; c.fill = hf; c.alignment = Alignment(horizontal="center")
        ws["A1"].alignment = Alignment(horizontal="left")
        for org in mat.index:
            ws.append([org] + [fmt(mat.loc[org, g]) for g in GENES]); r = ws.max_row
            for c in ws[r]:
                c.font = Font(name=F, bold=(org == SAL))
                if org == SAL: c.fill = salfill
            for col in range(2, 2 + len(GENES)): ws.cell(r, col).alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 20
        for i in range(len(GENES)): ws.column_dimensions[chr(66 + i)].width = 18
        ws.freeze_panes = "B2"

    sheet(wb.active, R, lambda v: round(float(v), 1) if pd.notna(v) else "")
    wb.active.title = "RNA nTPM"
    sheet(wb.create_sheet("Protein IHC"), P, lambda v: LVLNAME[int(v)] if pd.notna(v) else "n/a")
    out = f"{TAB}/Suppl_Table_HPA_expression.xlsx"; wb.save(out)
    R.to_csv(f"{TAB}/hpa_rna_nTPM_wide.csv"); P.applymap(lambda v: LVLNAME.get(v, "n/a") if pd.notna(v) else "n/a").to_csv(f"{TAB}/hpa_protein_IHC_wide.csv")
    print("wrote", out)


def figure(R, P):
    """Two-sided per-organ figure: RNA (left) and protein (right) on one shared viridis
    scale; RNA is normalized PER GENE (each gene to its own max organ) so a very high
    gene does not wash out the others. Salivary-gland row highlighted."""
    keep = (R.max(1) >= 1) | (P.max(1) >= 1)
    R, P = R[keep], P[keep]
    n, ng = len(R), len(GENES)
    gmax = {g: max(float(np.log10(R[g].fillna(0).values + 1.0).max()), 1e-9) for g in GENES}
    CMAP, BAND, MISS = plt.get_cmap("viridis"), "#F4A6A6", "#F2F2F2"
    LAB = {1: "L", 2: "M", 3: "H"}
    FS = dict(title=17, hdr=15, gene=13, org=12.5, tile=11, leg=12.5, cb=12)
    GAP = 1.5; x0p = ng + GAP; W = x0p + ng
    fig, ax = plt.subplots(figsize=(0.74 * 2 * ng + 6.5, 0.42 * n + 3.0))
    for i, org in enumerate(R.index):
        yy = n - 1 - i
        if org == SAL:
            ax.add_patch(Rectangle((-0.18, yy - 0.08), W + 0.36, 1.16, facecolor=BAND, ec="none", zorder=0))
        for j, g in enumerate(GENES):
            v = R.loc[org, g]
            rc = CMAP(np.log10((v if pd.notna(v) else 0) + 1.0) / gmax[g]) if pd.notna(v) else MISS
            ax.add_patch(Rectangle((j, yy), 1, 1, facecolor=rc, ec="white", lw=1.3, zorder=2))
            lv = P.loc[org, g]
            pc = CMAP(int(lv) / 3.0) if pd.notna(lv) else MISS
            ax.add_patch(Rectangle((x0p + j, yy), 1, 1, facecolor=pc, ec="white", lw=1.3, zorder=2))
            if pd.notna(lv) and lv >= 1:
                ax.text(x0p + j + 0.5, yy + 0.5, LAB[int(lv)], ha="center", va="center",
                        fontsize=FS["tile"], color="black" if lv >= 2 else "white", fontweight="bold", zorder=3)
    ax.set_xlim(-0.25, W + 0.25); ax.set_ylim(-0.25, n + 4.0)
    ax.set_yticks(np.arange(n) + 0.5); ax.set_yticklabels(R.index[::-1], fontsize=FS["org"])
    for t in ax.get_yticklabels():
        if t.get_text() == SAL: t.set_color("#B33A3A"); t.set_fontweight("bold")
    ax.set_xticks(list(np.arange(ng) + 0.5) + list(x0p + np.arange(ng) + 0.5))
    ax.set_xticklabels([f"$\\it{{{g}}}$" for g in GENES] * 2, fontsize=FS["gene"], rotation=90)
    ax.xaxis.set_ticks_position("top"); ax.tick_params(length=0)
    for s in ax.spines.values(): s.set_visible(False)
    ax.text(ng / 2, n + 3.3, "RNA expression (nTPM)", ha="center", fontsize=FS["hdr"], fontweight="bold")
    ax.text(x0p + ng / 2, n + 3.3, "Protein expression (IHC)", ha="center", fontsize=FS["hdr"], fontweight="bold")
    cax = fig.add_axes([0.13, -0.015, 0.22, 0.018])
    cb = fig.colorbar(ScalarMappable(norm=Normalize(0, 1), cmap=CMAP), cax=cax, orientation="horizontal")
    cb.set_ticks([0, 1]); cb.set_ticklabels(["min", "max"])
    cb.set_label("RNA (per gene, relative to its max organ)", fontsize=FS["cb"]); cb.ax.tick_params(labelsize=FS["cb"] - 1)
    leg = [Patch(fc=CMAP(i / 3.0), ec="white", label=LVLNAME[i]) for i in range(4)]
    ax.legend(handles=leg, title="Protein (IHC)", fontsize=FS["leg"], title_fontsize=FS["leg"],
              loc="upper left", bbox_to_anchor=(1.015, 1.0), frameon=False)
    fig.suptitle("Human Protein Atlas — RNA vs protein per organ, salivary glandular-backbone genes",
                 fontsize=FS["title"], fontweight="bold", y=1.045)
    fig.text(0.5, -0.055, "Data: Human Protein Atlas (proteinatlas.org); RNA consensus nTPM and antibody-based IHC. CC BY-SA 4.0.",
             ha="center", fontsize=9, color="#666")
    for ext in ("png", "svg", "pdf"):
        fig.savefig(f"{FIG}/HPA_rna_protein_combined.{ext}", dpi=200 if ext == "png" else None, bbox_inches="tight", facecolor="white")
    print("wrote", f"{FIG}/HPA_rna_protein_combined.[png|svg|pdf]")


if __name__ == "__main__":
    rna, pro = fetch()
    R, P = matrices(rna, pro)
    write_table(R, P)
    figure(R, P)
